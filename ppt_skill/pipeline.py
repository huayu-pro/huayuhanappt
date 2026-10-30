"""生成流程编排与内容校验。

校验从"打印几行文字 + sys.exit(1)"升级为结构化报告：
每条 ValidationError 都带 page_num / field / message / suggestion，
告诉用户"改哪一页的哪个字段、怎么改"，而不是甩一个页码列表让人猜。
"""

import time
from dataclasses import dataclass, field
from pathlib import Path
from typing import List, Optional

from pptx import Presentation

from .config import LIMITS
from .errors import RenderError, TimeoutError
from .models import SlideContent, VALID_CHART_TYPES, VALID_SLIDE_TYPES
from .rendering.slides import SlideRenderer


@dataclass
class ValidationError:
    page_num: Optional[int]
    field: str
    message: str
    suggestion: str

    def render(self):
        where = f"第 {self.page_num} 页" if self.page_num else "整体"
        return f"  • [{where}] {self.field}：{self.message}\n    → {self.suggestion}"


@dataclass
class GenerationReport:
    output_path: Optional[Path] = None
    total_pages: int = 0
    theme: str = ""
    images_used: int = 0
    image_failures: list = field(default_factory=list)
    warnings: List[str] = field(default_factory=list)
    validation_errors: List[ValidationError] = field(default_factory=list)
    elapsed_seconds: float = 0.0
    skipped_images: bool = False

    @property
    def ok(self):
        return self.output_path is not None and not self.validation_errors


def validate_content(slides: List[SlideContent]) -> List[ValidationError]:
    """内容校验，返回结构化错误列表（空列表表示通过）。"""
    errors: List[ValidationError] = []
    total = len(slides)
    min_pages = LIMITS["min_pages"]
    max_pages = LIMITS["max_pages"]

    if total < min_pages or total > max_pages:
        hint = (f"增减内容使总页数落在 {min_pages}-{max_pages} 页；学术汇报推荐 "
                f"{LIMITS['recommended_pages'][0]}-{LIMITS['recommended_pages'][1]} 页")
        if total < min_pages:
            hint += f"。文稿太短时脚本会自动拆分段落补齐，仍不足 {min_pages} 页说明内容量不够——请补充研究背景、方法或结论等内容"
        errors.append(ValidationError(None, "页数", f"当前 {total} 页，超出 {min_pages}-{max_pages} 页范围", hint))

    seen_pages = {}
    for index, slide in enumerate(slides):
        prefix = f"第 {index + 1} 项"
        if slide.page_num in seen_pages:
            errors.append(ValidationError(
                slide.page_num, "page_num",
                f"页码重复（与{seen_pages[slide.page_num]}相同）",
                "把每一页的 page_num 改成唯一整数；或删掉 page_num 让脚本自动编号",
            ))
        seen_pages[slide.page_num] = prefix

        if slide.slide_type not in VALID_SLIDE_TYPES:
            errors.append(ValidationError(
                slide.page_num, "slide_type",
                f"取值 {slide.slide_type!r} 不合法",
                f"只能是 {' / '.join(VALID_SLIDE_TYPES)}",
            ))

        if slide.chart_type not in VALID_CHART_TYPES:
            errors.append(ValidationError(
                slide.page_num, "chart_type",
                f"取值 {slide.chart_type!r} 不合法",
                "只能是 pie / bar / line，或留空让脚本按数据自动推断",
            ))

        if slide.chart_type and not slide.table_data:
            errors.append(ValidationError(
                slide.page_num, "table_data",
                f"指定了 chart_type={slide.chart_type!r} 但没有表格数据",
                "补上 table_data（headers + rows），或把 chart_type 留空",
            ))

        if slide.table_data and not slide.table_data.rows:
            errors.append(ValidationError(
                slide.page_num, "table_data",
                "rows 是空的",
                "至少写一行数据，否则这一页不出图表",
            ))

        if len(slide.bullet_points) > LIMITS["max_bullets_per_slide"]:
            errors.append(ValidationError(
                slide.page_num, "bullet_points",
                f"要点 {len(slide.bullet_points)} 条，超过单页上限 "
                f"{LIMITS['max_bullets_per_slide']} 条",
                "拆分到两页，或合并相近要点；单页要点过多会溢出版面",
            ))

        if slide.slide_type == "content" and not (
                slide.bullet_points or slide.quote_text
                or slide.code_block or slide.table_data):
            errors.append(ValidationError(
                slide.page_num, "内容",
                "正文页没有任何内容元素",
                "至少填 bullet_points / quote_text / code_block / table_data 之一",
            ))

    title_pages = [s.page_num for s in slides if s.slide_type == "title"]
    if not title_pages:
        errors.append(ValidationError(
            None, "封面页", "缺少 slide_type='title' 的页面",
            "把第一页改成 {\"slide_type\": \"title\", \"title\": \"你的题目\"}",
        ))
    elif len(title_pages) > 1:
        errors.append(ValidationError(
            None, "封面页", f"slide_type='title' 出现 {len(title_pages)} 次"
            f"（第 {title_pages} 页）",
            "一份 PPT 只保留一个封面页，多余的改成 content 或 section",
        ))
    thank_pages = [s.page_num for s in slides if s.slide_type == "thank"]
    if not thank_pages:
        errors.append(ValidationError(
            None, "致谢页", "缺少 slide_type='thank' 的页面",
            "在末尾加一页 {\"slide_type\": \"thank\", \"title\": \"感谢聆听\"}",
        ))
    elif len(thank_pages) > 1:
        errors.append(ValidationError(
            None, "致谢页", f"slide_type='thank' 出现 {len(thank_pages)} 次"
            f"（第 {thank_pages} 页）",
            "只保留末尾一页致谢页",
        ))

    numbers = sorted(s.page_num for s in slides)
    if numbers != list(range(1, total + 1)):
        errors.append(ValidationError(
            None, "page_num",
            f"页码不是 1 到 {total} 的连续整数（当前：{numbers[:8]}"
            f"{'…' if len(numbers) > 8 else ''}）",
            "让 page_num 从 1 开始连续编号；最简单的方法是全部省略 page_num，由脚本自动编号",
        ))

    return errors


def generate_ppt(slides, config, meta=None, image_map=None, image_failures=None,
                 timeout=1800):
    """完整生成流程：校验 → 建目录 → 渲染 → 保存。

    image_map: {page_num: [Path, ...]}，由调用方（pipeline 或 UI）预先爬取。
    image_failures: [dict, ...] 爬取失败明细，原样带进报告供末尾汇总。
    返回 GenerationReport，绝不静默失败。
    """
    started = time.time()
    report = GenerationReport(theme=config.theme, total_pages=len(slides))
    image_map = image_map or {}
    report.image_failures = list(image_failures or [])

    # 1. 内容校验（启动期，失败就不进入渲染）
    report.validation_errors = validate_content(slides)
    if report.validation_errors:
        return report

    # 2. 输出路径早期校验（原版在 30 页渲染完后才因文件名崩溃）
    output_path = config.resolve_output()

    # 3. 渲染
    try:
        presentation = Presentation()
        presentation.slide_width = SlideRenderer.SLIDE_WIDTH
        presentation.slide_height = SlideRenderer.SLIDE_HEIGHT

        renderer = SlideRenderer(
            presentation, config.palette, config.theme,
            config.chart_colors, meta=meta)

        total = len(slides)
        for slide in slides:
            # 在渲染循环内检查超时：原版在 save() 之后才判断，
            # 超时抛出时文件已经落盘，用户以为失败其实产物是好的。
            if time.time() - started > timeout:
                raise TimeoutError(
                    f"渲染第 {slide.page_num} 页时已超过 {timeout} 秒上限",
                    "减少页数或关闭自动配图（--skip-images）后重试",
                )
            paths = image_map.get(slide.page_num) or []
            image_path = paths[0] if paths else None
            renderer.add_slide(slide, image_path, total)
            # 用户显式 --skip-images 时不提示"期望配图"——那是选择不是失败
            if (not config.skip_images
                    and slide.slide_type in ("title", "section", "thank")
                    and slide.image_keywords and not image_path):
                report.warnings.append(
                    f"第 {slide.page_num} 页期望配图但未获取到，已按无图版式渲染")

        if time.time() - started > timeout:
            raise TimeoutError(
                f"生成耗时超过 {timeout} 秒上限",
                "减少页数或关闭自动配图（--skip-images）后重试",
            )
        presentation.save(str(output_path))
    except TimeoutError:
        raise
    except Exception as exc:
        raise RenderError(
            f"PPT 渲染或保存失败：{type(exc).__name__}: {exc}",
            "检查输出目录权限与磁盘空间；若刚改过内容文件，先用校验命令排查格式",
        )

    try:
        _normalize_pkg_timestamps(output_path)
    except Exception as exc:
        # 归一化失败不影响产物可用性，只损失字节级可复现
        print(f"  ⚠️  产物时间戳归一化跳过：{exc}")

    report.output_path = output_path
    report.warnings.extend(renderer.warnings)
    report.images_used = sum(1 for paths in image_map.values() for _ in paths)
    report.elapsed_seconds = round(time.time() - started, 2)
    return report




def _normalize_pkg_timestamps(path, fixed_date=(2026, 1, 1, 0, 0, 0)):
    """把 pptx 内所有 zip 条目与嵌入工作簿的时间戳固定为常量。

    python-pptx 与内嵌 xlsxwriter 会把当前时间写进 zip 头与
    docProps/core.xml——同一输入两次运行产出字节不同的文件，
    且把生成时刻泄露进交付物。统一改写成固定时间，使产物字节级可复现。
    """
    import io
    import os
    import re
    import zipfile

    src = Path(path)
    tmp = src.with_suffix(".tmp.pptx")
    created_re = re.compile(rb"(<dcterms:created[^>]*>)[^<]*(</dcterms:created>)")
    modified_re = re.compile(rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)")
    fixed_xml = b"2026-01-01T00:00:00Z"

    with zipfile.ZipFile(src, "r") as zin, \
            zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data = zin.read(info.filename)
            if (info.filename.startswith("ppt/embeddings/")
                    and info.filename.endswith(".xlsx")):
                inner = io.BytesIO()
                with zipfile.ZipFile(io.BytesIO(data), "r") as x_in, \
                        zipfile.ZipFile(inner, "w", zipfile.ZIP_DEFLATED) as x_out:
                    for x_info in x_in.infolist():
                        blob = x_in.read(x_info.filename)
                        if x_info.filename == "docProps/core.xml":
                            blob = created_re.sub(rb"\g<1>" + fixed_xml + rb"\g<2>", blob)
                            blob = modified_re.sub(rb"\g<1>" + fixed_xml + rb"\g<2>", blob)
                        out_info = zipfile.ZipInfo(x_info.filename, date_time=fixed_date)
                        out_info.compress_type = zipfile.ZIP_DEFLATED
                        out_info.external_attr = x_info.external_attr
                        x_out.writestr(out_info, blob)
                data = inner.getvalue()
            out_info = zipfile.ZipInfo(info.filename, date_time=fixed_date)
            out_info.compress_type = zipfile.ZIP_DEFLATED
            out_info.external_attr = info.external_attr
            zout.writestr(out_info, data)
    os.replace(tmp, src)
