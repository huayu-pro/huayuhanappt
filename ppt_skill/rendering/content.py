"""内容渲染：列表、有序列表、引用块、代码块、表格。

原版这三个函数是死代码，从未被调用：
    _render_ordered_list（宣传"编号圆圈"）
    _render_table（宣传"表头底色+交替行"）
    detect_chart_type_from_table（宣传"自动检测图表类型"）
v2.0.0 把它们全部接进调用链，SKILL.md 的宣传与实现对齐。
"""

from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Pt

from ..text_utils import estimate_lines
from .base import to_rgb


class ContentRenderer:
    """各类内容元素的渲染。"""

    def __init__(self, base: "Renderer"):
        self.base = base

    # ---- 无序列表 ----
    def render_bullet_list(self, slide, items, left, top, width, height,
                           font_size=Pt(15), line_spacing=1.6):
        """主题色实心圆点 + 悬挂缩进（圆点由 buChar 提供，文本不再手写 ●）。

        按内容预算高度：单条要点折行后超出剩余空间即截断，
        返回实际渲染条数——调用方据此告警，而不是任由文字画出版面。
        空列表返回 0。
        """
        items = [str(i).strip() for i in items if str(i).strip()]
        if not items:
            return 0

        # 每行可容纳的半角列数（全角按 2 列），再换算成行数预算
        chars_per_line = max(8, int(width / 914400 * 72 / font_size.pt * 1.9))
        line_height = Emu(int(font_size.pt * line_spacing * 12700))

        box = slide.shapes.add_textbox(left, top, width, height)
        tf = box.text_frame
        tf.word_wrap = True

        rendered = 0
        used = Emu(0)
        for item in items:
            lines = estimate_lines(item, chars_per_line)
            needed = line_height * lines + Emu(int(font_size.pt * 0.4 * 12700))
            if used + needed > height and rendered > 0:
                break
            p = tf.paragraphs[0] if rendered == 0 else tf.add_paragraph()
            p.text = item
            p.font.size = font_size
            p.font.color.rgb = self.base.color("text")
            p.font.name = "Arial"
            p.space_after = Pt(font_size.pt * 0.4)
            self.base.set_bullet_color(p, self.base.color("accent"))
            used += needed
            rendered += 1
        return rendered

    # ---- 有序列表 ----
    def render_ordered_list(self, slide, items, left, top, width, height,
                            font_size=Pt(15)):
        """主题色编号圆圈 + 缩进文字。返回实际渲染条数（0 = 没画）。"""
        items = [i for i in items if str(i).strip()]
        if not items:
            return 0

        y = top
        rendered = 0
        for index, item in enumerate(items):
            if y + Emu(480000) > top + height:
                break
            circle = slide.shapes.add_shape(
                MSO_SHAPE.OVAL, left, y + Emu(60000),
                Emu(250000), Emu(250000))
            circle.fill.solid()
            circle.fill.fore_color.rgb = self.base.color("accent")
            circle.line.fill.background()
            circle.shadow.inherit = False

            tf = circle.text_frame
            tf.word_wrap = False
            p = tf.paragraphs[0]
            p.text = str(index + 1)
            p.font.size = Pt(10)
            p.font.color.rgb = to_rgb((0xFF, 0xFF, 0xFF))
            p.font.bold = True
            p.font.name = "Arial"
            p.alignment = PP_ALIGN.CENTER

            self.base.add_textbox(
                slide, left + Emu(350000), y,
                width - Emu(450000), Emu(400000),
                item, font_size=font_size)
            y += Emu(480000)
            rendered += 1
        return rendered

    # ---- 引用块 ----
    def render_quote(self, slide, text, left, top, width, max_height=None):
        """左侧主题色竖条 + 浅色背景 + 斜体。

        高度按"每行可容纳字数"估算，考虑自动换行——
        原版只数 \n 个数，一个 200 字段落按 1 行算，文字直接溢出底色。
        max_height 给出剩余可用空间（如页码区上沿），触顶时收敛高度。
        """
        if not str(text).strip():
            return Emu(0)

        chars_per_line = max(10, int(width / 914400 * 72 / 13 * 2.2))
        lines = estimate_lines(text, chars_per_line)
        cap = 2280000 if max_height is None else min(2280000, int(max_height))
        if cap < 550000:
            return Emu(0)
        block_height = Emu(int(min(cap, max(550000, lines * 280000 + 250000))))

        bar = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, left, top, Emu(57000), block_height)
        bar.fill.solid()
        bar.fill.fore_color.rgb = self.base.color("accent")
        bar.line.fill.background()
        bar.shadow.inherit = False

        bg = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, left + Emu(57000), top,
            width - Emu(57000), block_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = self.base.color("light")
        bg.line.fill.background()
        bg.shadow.inherit = False

        self.base.add_textbox(
            slide, left + Emu(228000), top + Emu(57000),
            width - Emu(342000), block_height - Emu(114000),
            text, font_size=Pt(13),
            font_color=self.base.color("text"), italic=True)
        return block_height

    # ---- 代码块 ----
    def render_code_block(self, slide, code, left, top, width, max_height=None):
        """深色背景 + Consolas 等宽 + 圆角。空代码返回 0 高度。

        max_height 给出剩余可用空间，触顶时收敛高度，避免代码块
        被推出幻灯片底边（长引用 + 长代码叠加时会出现）。
        """
        if not str(code).strip():
            return Emu(0)

        chars_per_line = max(20, int(width / 914400 * 72 / 10 * 1.6))
        lines = estimate_lines(code, chars_per_line)
        cap = 2850000 if max_height is None else min(2850000, int(max_height))
        if cap < 350000:
            return Emu(0)
        block_height = Emu(int(min(cap, max(350000, lines * 250000 + 350000))))

        bg = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, left - Emu(114000), top,
            width + Emu(228000), block_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = to_rgb((0x1E, 0x1E, 0x1E))
        bg.line.fill.background()
        bg.shadow.inherit = False
        try:
            bg.adjustments[0] = 0.05
        except Exception:
            pass

        self.base.add_textbox(
            slide, left + Emu(57000), top + Emu(114000),
            width, block_height - Emu(228000),
            code, font_size=Pt(10),
            font_color=to_rgb((0xD4, 0xD4, 0xD4)), font_name="Consolas")
        return block_height

    # ---- 表格 ----
    def render_table(self, slide, headers, rows, left, top, width, max_height=None,
                     page_num=None):
        """表头主题色底 + 白字，交替行浅色着色。

        原版这是死代码；v2.0.0 起它同时是图表失败时的降级路径。
        超出 max_height 的行会被截断并记 warning——原实现静默丢行，
        100 行的表只画 9 行用户毫不知情。
        """
        if not rows:
            return Emu(0)
        if max_height is None:
            max_height = Emu(4560000)

        col_count = max([len(headers)] + [len(r) for r in rows]) if rows else len(headers)
        col_count = max(1, col_count)
        col_width = width // col_count
        row_height = Emu(400000)

        y = top
        rendered_rows = 0
        for r in range(1 + len(rows)):
            if y + row_height > top + max_height:
                break
            if r == 0:
                self.base.add_shape(
                    slide, MSO_SHAPE.RECTANGLE,
                    left - Emu(57000), y,
                    col_count * col_width + Emu(114000), row_height,
                    fill_color=self.base.color("primary"))
            elif r % 2 == 0:
                self.base.add_shape(
                    slide, MSO_SHAPE.RECTANGLE,
                    left - Emu(57000), y,
                    col_count * col_width + Emu(114000), row_height,
                    fill_color=self.base.color("lighter"))

            for c in range(col_count):
                if r == 0:
                    cell = headers[c] if c < len(headers) else ""
                else:
                    cell = rows[r - 1][c] if c < len(rows[r - 1]) else ""
                is_header = (r == 0)
                self.base.add_textbox(
                    slide, left + c * col_width, y,
                    col_width - Emu(57000), row_height,
                    cell,
                    font_size=Pt(11) if is_header else Pt(10),
                    font_color=to_rgb((0xFF, 0xFF, 0xFF)) if is_header else self.base.color("text"),
                    bold=is_header,
                    alignment=PP_ALIGN.CENTER if is_header else PP_ALIGN.LEFT)
            y += row_height
            rendered_rows += 1

        dropped = len(rows) - (rendered_rows - 1) if rendered_rows else len(rows)
        if dropped > 0:
            where = f"第 {page_num} 页" if page_num else "当前页"
            self.base.warnings.append(
                f"{where}表格空间不足，仅展示前 {rendered_rows - 1} 行（共 {len(rows)} 行，"
                f"截断 {dropped} 行）")
        return y - top
