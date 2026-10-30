"""图表生成：饼图 / 柱状图 / 折线图，含真实降级链。

相对原版的四点修复：

H2 饼图多系列静默丢数据：python-pptx 对 PIE 只保留第一个系列
   （实测 len(chart.series)==1）。这里在传数据前就把多系列饼图
   截断为首系列，并打印警告，让用户知道发生了什么。
H3 表头语义：labels 取每一行的第一列，不再把第一行数据当表头吞掉。
F15 负值与缺失值：饼图遇负值降级为柱状图；解析不出的单元格让整列
   放弃（原版把 None 替换成 0，把"缺失"伪装成"真的是 0"）。
M2 降级失效：图表生成失败时返回 False，由 slides.py 调用 render_table
    兜底，而不是只留一个光秃秃的标题。
"""

from pptx.chart.data import CategoryChartData
from pptx.enum.chart import (XL_CHART_TYPE, XL_LABEL_POSITION,
                             XL_LEGEND_POSITION)
from pptx.util import Pt

from ..text_utils import _numeric_columns, table_labels
from .base import hex_to_rgb


class ChartRenderer:
    """把 TableData 渲染成图表。"""

    def __init__(self, base: "Renderer"):
        self.base = base

    def add_chart(self, slide, chart_type, table, left, top, width, height):
        """返回 True 表示图表已渲染；False 表示调用方应走降级路径。"""
        if not table or not table.rows:
            return False

        columns = _numeric_columns(table)
        if not columns:
            print("    ⚠️  表格里没有完整可解析的数值列，改用表格展示")
            return False

        labels = table_labels(table)
        colors = self.base.chart_colors

        # 饼图只接受单系列，且不能有负值
        if chart_type == "pie":
            if len(columns) > 1:
                dropped = "、".join(name for name, _ in columns[1:])
                print(f"    ⚠️  饼图只展示一个数值系列，已忽略：{dropped}")
            columns = columns[:1]
            if any(v < 0 for v in columns[0][1]):
                print("    ⚠️  饼图不支持负值，已改用柱状图")
                chart_type = "bar"

        chart_data = CategoryChartData()
        chart_data.categories = labels
        for name, values in columns:
            chart_data.add_series(name, values)

        xl_type = {
            "pie": XL_CHART_TYPE.PIE,
            "line": XL_CHART_TYPE.LINE,
            "bar": XL_CHART_TYPE.BAR_CLUSTERED,
        }.get(chart_type, XL_CHART_TYPE.BAR_CLUSTERED)

        try:
            frame = slide.shapes.add_chart(xl_type, left, top, width, height, chart_data)
        except Exception as exc:
            print(f"    ⚠️  图表生成失败，改用表格展示：{exc}")
            return False

        self._style_chart(frame.chart, chart_type, colors, len(columns))
        return True

    def _style_chart(self, chart, chart_type, colors, series_count):
        chart.has_legend = True
        try:
            if chart_type == "pie":
                chart.legend.position = XL_LEGEND_POSITION.RIGHT
            else:
                chart.legend.position = XL_LEGEND_POSITION.BOTTOM
            chart.legend.include_in_layout = False
        except Exception:
            pass

        try:
            plot = chart.plots[0]
            plot.has_data_labels = True
            labels = plot.data_labels
            if chart_type == "pie":
                labels.show_percentage = True
                labels.show_value = False
                labels.font.size = Pt(10)
            else:
                labels.show_value = True
                labels.font.size = Pt(9)
                if chart_type == "bar":
                    try:
                        labels.position = XL_LABEL_POSITION.OUTSIDE_END
                    except Exception:
                        pass
            labels.font.color.rgb = self.base.color("text")
        except Exception:
            pass

        # 系列配色：pie 只有一个系列时逐点染，多系列按系列染
        try:
            for index, series in enumerate(chart.series):
                series.format.fill.solid()
                if chart_type == "pie" and series_count == 1:
                    continue
                color = colors[index % len(colors)]
                series.format.fill.fore_color.rgb = hex_to_rgb(color)
            if chart_type == "pie":
                self._color_pie_points(chart, colors)
        except Exception:
            pass

    def _color_pie_points(self, chart, colors):
        """饼图按数据点逐个上色（单系列时系列级上色不生效）。"""
        try:
            series = chart.series[0]
            for index, point in enumerate(series.points):
                color = colors[index % len(colors)]
                point.format.fill.solid()
                point.format.fill.fore_color.rgb = hex_to_rgb(color)
        except Exception:
            pass
